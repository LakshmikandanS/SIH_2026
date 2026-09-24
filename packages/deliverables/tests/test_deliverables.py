"""Deliverables: normalisation, grounding, rendering into the fixed templates, the
four-tier ladder, and approval to a frozen release.

The integration tests run against a real Postgres (the same scratch database the
platform tests use) with a document, its pages and task-scoped evidence written the
way citadel_knowledge writes them, so tier 3 resolves citations against real rows.
"""

from __future__ import annotations

import hashlib
import io
from pathlib import Path
from typing import Any, Iterator

import pytest

from citadel_contracts.domain import User
from citadel_deliverables import (
    ArtifactError,
    TaskFacts,
    cited_ids,
    decide,
    generate,
    get_artifact,
    is_grounded,
    normalise,
    quantities,
    read_bytes,
    rendered_text,
    to_html,
)
from citadel_deliverables.content import Item
from citadel_deliverables.render import arrange_cells
from citadel_platform.db import Database, Json, RealArray
from citadel_platform.registry.loader import load_templates
from citadel_platform.registry.schema import TemplateEntry
from citadel_platform.storage import DataDir
from pg_scratch import apply_all_migrations, pg_scratch_db, requires_pgvector

REPO_ROOT = Path(__file__).resolve().parents[3]
REGISTRY_DIR = REPO_ROOT / "registry"
ENGINEER = User("demo-engineer-1", "R. Kulkarni", ("engineer",), "internal", "process-engineering")
APPROVER = User("demo-approver", "A. Menon", ("approver",), "confidential", "quality-assurance")


def _templates() -> dict[str, TemplateEntry]:
    return {t.id: t for t in load_templates(REGISTRY_DIR)}


GOOD_NOTE: dict[str, Any] = {
    "subject": "Continued service of heat exchanger E-101 pending shell repair",
    "background": "Heat exchanger E-101 was inspected under IR-2026-0147 after the flange leak incident.",
    "findings": [
        {"text": "CML-3 measured a minimum wall thickness of 9.2 mm against a t-min of 8.4 mm.", "citations": ["E1"]},
        "The corrosion rate at CML-3 is 0.25 mm/year [E1]",
    ],
    "analysis": {"text": "At 0.25 mm/year the remaining life to t-min is 3.2 years.", "citations": ["E1", "E2"]},
    "recommendation": "Continue service with a re-inspection of CML-3 within 12 months.",
    "annexures": ["Inspection report IR-2026-0147 [E1]"],
}


# -- pure: normalisation and grounding ----------------------------------------------------


def test_normalise_reduces_loose_content_to_one_shape_and_records_mismatches():
    template = _templates()["approval-note"]
    content = normalise(template, {**GOOD_NOTE, "surprise": "x", "findings": "just one finding [E2]"})
    assert content.sections["findings"].items[0].text == "just one finding"
    assert content.sections["findings"].items[0].citations == ["E2"]
    assert any("'surprise' is not a section" in i for i in content.issues)
    assert any("expected a list" in i for i in content.issues)
    assert cited_ids(template, GOOD_NOTE) == ["E1", "E2"]


def test_a_missing_required_section_is_a_schema_issue_not_a_guess():
    template = _templates()["approval-note"]
    content = normalise(template, {k: v for k, v in GOOD_NOTE.items() if k != "findings"})
    assert content.sections["findings"].empty
    assert any("findings" in i and "at least 1" in i for i in content.issues)


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("CML-3 on E-101 reads 9.2 mm", ["9.2"]),
        ("per IR-2026-0147, page 2", ["2"]),
        ("inspected on 2026-03-14 and 14 March 2026", []),
        ("0.25 mm/year gives 3.2 years [E1]", ["0.25", "3.2"]),
    ],
)
def test_quantities_ignores_identifiers_dates_and_citation_markers(text: str, expected: list[str]):
    assert quantities(text) == expected


def test_grounding_tolerates_rounding_down_never_invented_precision():
    assert is_grounded("3.2", ["3.24"])  # a rounded source value is the same claim
    assert not is_grounded("3.24", ["3.2"])  # more precision than the source is invention
    assert is_grounded("9.20", ["9.2"])
    assert not is_grounded("9.3", ["9.2"])
    assert is_grounded("0.801", ["0.8"], relative_tolerance=0.002)


@requires_pgvector
@pytest.mark.integration
def test_references_sit_before_the_full_stop_and_short_items_still_verify(world: dict[str, Any]):
    db, data_dir = world["db"], world["data_dir"]
    template = _templates()["approval-note"]
    task = _task(db)
    evidence = _evidence(db, task, world["document_id"])
    content = {**GOOD_NOTE, "findings": ["Wall loss at CML-3 [E1]."], "background": "Inspected under IR-2026-0147 [E1] ."}
    generated = generate(db, data_dir, REGISTRY_DIR, template, content, task=task, author=ENGINEER, evidence=evidence)
    assert generated.verification.passed, generated.verification.to_dict()
    _, data = read_bytes(db, data_dir, generated.artifact_id)
    text = rendered_text("docx", data)
    assert "Wall loss at CML-3 [1]." in text and "Inspected under IR-2026-0147 [1]." in text


def test_table_rows_follow_the_templates_own_column_headings():
    item = Item(text="", citations=["E2"], cells={"measurement": "9.2 mm", "location": "CML-3", "defect": "wall loss"})
    row = arrange_cells(["Location", "Defect", "Measurement", "Ref"], item, {"E2": 1})
    assert row == ["CML-3", "wall loss", "9.2 mm", "[1]"]
    # nothing is dropped: an unmatched value joins the last content column
    extra = Item(text="", cells={"location": "CML-4", "note": "near weld"})
    assert arrange_cells(["Location", "Ref"], extra, {}) == ["CML-4; near weld", ""]


# -- integration: real rows, real files ----------------------------------------------------


@pytest.fixture(scope="module")
def world(tmp_path_factory: pytest.TempPathFactory) -> Iterator[dict[str, Any]]:
    with pg_scratch_db() as env:
        apply_all_migrations(env)
        db = Database(env=env)
        data_dir = DataDir(root=tmp_path_factory.mktemp("data"))
        document_id = "11111111-1111-4111-8111-111111111111"
        db.script([
            (
                "INSERT INTO documents (id, title, filename, mime_type, classification, acl, uploaded_by, status, version, "
                "sha256, storage_ref, page_count) VALUES (%(id)s::uuid, 'Inspection report IR-2026-0147 (E-101)', "
                "'ir.pdf', 'application/pdf', 'internal', ARRAY['process-engineering'], 'demo-engineer-1', 'ready', 1, "
                "%(sha)s, 'documents/x/ir.pdf', 2)",
                {"id": document_id, "sha": "0" * 64},
            ),
            ("INSERT INTO document_pages (document_id, version, page, width_px, height_px, text_source) VALUES "
             "(%(id)s::uuid, 1, 1, 1240, 1754, 'ocr'), (%(id)s::uuid, 1, 2, 1240, 1754, 'ocr')", {"id": document_id}),
        ])
        yield {"env": env, "db": db, "data_dir": data_dir, "document_id": document_id}


def _task(db: Database, classification: str = "internal") -> TaskFacts:
    task_id = str(db.scalar(
        "INSERT INTO tasks (goal, submitted_by, classification) VALUES ('Draft an approval note for E-101', "
        "(SELECT id FROM users WHERE external_identity = 'demo-engineer-1'), %(c)s) RETURNING id::text",
        {"c": classification},
    ))
    return TaskFacts(task_id, classification.upper(), "Draft an approval note for E-101")


def _evidence(db: Database, task: TaskFacts, document_id: str, *, classification: str = "internal") -> dict[str, Any]:
    rows = [
        ("E1", 2, [0.08, 0.30, 0.92, 0.46],
         "CML-3 minimum measured thickness 9.2 mm; t-min 8.4 mm; corrosion rate 0.25 mm/year; remaining life 3.2 years"),
        ("E2", 1, [0.08, 0.12, 0.92, 0.20], "Heat exchanger E-101 shell, carbon steel, nominal 12.0 mm"),
    ]
    db.script([
        (
            "INSERT INTO task_evidence (task_id, evidence_id, kind, document_id, version, page, bbox, text, classification, detail) "
            "VALUES (%(t)s::uuid, %(e)s, 'document', %(d)s::uuid, 1, %(p)s, %(b)s, %(x)s, %(c)s, %(detail)s)",
            {"t": task.task_id, "e": e, "d": document_id, "p": p, "b": RealArray(b), "x": x, "c": classification,
             "detail": Json({"title": "Inspection report IR-2026-0147 (E-101)"})},
        )
        for e, p, b, x in rows
    ])
    resolved = db.query(
        "SELECT e.evidence_id, e.kind, e.document_id::text AS document_id, e.version, e.page, e.bbox, e.text, "
        "e.classification, e.detail, d.title FROM task_evidence e LEFT JOIN documents d ON d.id = e.document_id "
        "WHERE e.task_id = %(t)s::uuid",
        {"t": task.task_id},
    )
    return {r["evidence_id"]: r for r in resolved}


@requires_pgvector
@pytest.mark.integration
def test_an_approval_note_is_rendered_verified_approved_and_released(world: dict[str, Any]):
    db, data_dir = world["db"], world["data_dir"]
    template = _templates()["approval-note"]
    task = _task(db)
    evidence = _evidence(db, task, world["document_id"])

    generated = generate(db, data_dir, REGISTRY_DIR, template, GOOD_NOTE, task=task, author=ENGINEER, evidence=evidence)
    tiers = {t.name: t.status for t in generated.verification.tiers}
    assert tiers == {"structural": "pass", "schema": "pass", "citation": "pass", "grounding": "pass"}, generated.verification.to_dict()
    assert generated.status == "VERIFIED"

    artifact, data = read_bytes(db, data_dir, generated.artifact_id)
    text = rendered_text("docx", data)
    assert text.count("INTERNAL") >= 3  # header, banner, footer
    assert "9.2 mm" in text and "[1]" in text and "PENDING APPROVAL" in text
    assert "Inspection report IR-2026-0147 (E-101)" in text  # the sources table
    assert "(0.08, 0.30, 0.92, 0.46)" in text  # the exact region
    assert "{{" not in text
    html = to_html("docx", data)
    assert "<table>" in html and "INTERNAL" in html

    with pytest.raises(ArtifactError, match="cannot approve"):
        decide(db, data_dir, REGISTRY_DIR, list(_templates().values()), artifact_id=generated.artifact_id,
               approver=ENGINEER, approve=True, comment="self-approval")

    result = decide(db, data_dir, REGISTRY_DIR, list(_templates().values()), artifact_id=generated.artifact_id,
                    approver=APPROVER, approve=True, comment="Agree with the re-inspection interval.")
    assert result["status"] == "RELEASED"
    released, released_bytes = read_bytes(db, data_dir, generated.artifact_id)
    assert released["status"] == "RELEASED"
    assert hashlib.sha256(released_bytes).hexdigest() == result["sha256"] != generated.sha256
    released_text = rendered_text("docx", released_bytes)
    assert "A. Menon (quality-assurance)" in released_text and "PENDING APPROVAL" not in released_text
    assert "Approved for release" in released_text
    record = released["provenance"]["record"]
    assert record["task"]["submitted_by"] == "demo-engineer-1"
    assert [e["evidence_id"] for e in record["evidence"]] == ["E1", "E2"]
    assert record["approvals"][0]["approver"] == "demo-approver"
    assert record["release"]["draft_sha256"] == generated.sha256

    # frozen: the database refuses any further change to a RELEASED row
    with pytest.raises(Exception, match="immutable"):
        db.execute("UPDATE artifacts SET title = 'edited' WHERE id = %(id)s::uuid", {"id": generated.artifact_id})
    with pytest.raises(ArtifactError):
        decide(db, data_dir, REGISTRY_DIR, list(_templates().values()), artifact_id=generated.artifact_id,
               approver=APPROVER, approve=True, comment="again")


@requires_pgvector
@pytest.mark.integration
def test_an_untraceable_number_is_flagged_not_dropped_and_not_passed(world: dict[str, Any]):
    db, data_dir = world["db"], world["data_dir"]
    template = _templates()["approval-note"]
    task = _task(db)
    evidence = _evidence(db, task, world["document_id"])
    content = {**GOOD_NOTE, "findings": [{"text": "CML-3 measured 7.9 mm at the nozzle.", "citations": ["E1"]}]}
    generated = generate(db, data_dir, REGISTRY_DIR, template, content, task=task, author=ENGINEER, evidence=evidence)
    grounding = generated.verification.tiers[3]
    assert grounding.status == "flagged"
    assert generated.verification.flagged_claims[0]["number"] == "7.9"
    assert generated.status == "VERIFIED"  # tier 4 travels with the artifact to the approver


@requires_pgvector
@pytest.mark.integration
def test_a_citation_the_task_was_never_given_fails_tier_three(world: dict[str, Any]):
    db, data_dir = world["db"], world["data_dir"]
    template = _templates()["approval-note"]
    task = _task(db)
    evidence = _evidence(db, task, world["document_id"])
    content = {**GOOD_NOTE, "annexures": ["A document nobody retrieved [E9]"]}
    generated = generate(db, data_dir, REGISTRY_DIR, template, content, task=task, author=ENGINEER, evidence=evidence)
    citation = generated.verification.tiers[2]
    assert citation.status == "fail" and any("E9" in i for i in citation.issues)
    assert generated.status == "TEMP"
    assert get_artifact(db, generated.artifact_id)["status"] == "TEMP"  # type: ignore[index]


@requires_pgvector
@pytest.mark.integration
def test_evidence_above_the_documents_marking_fails_tier_three(world: dict[str, Any]):
    db, data_dir = world["db"], world["data_dir"]
    template = _templates()["approval-note"]
    task = _task(db, "internal")
    evidence = _evidence(db, task, world["document_id"], classification="confidential")
    generated = generate(db, data_dir, REGISTRY_DIR, template, GOOD_NOTE, task=task, author=ENGINEER, evidence=evidence)
    assert generated.verification.tiers[2].status == "fail"
    assert any("CONFIDENTIAL evidence in a document marked INTERNAL" in i for i in generated.verification.tiers[2].issues)


@requires_pgvector
@pytest.mark.integration
def test_a_rejection_keeps_the_draft_and_the_revision_carries_the_history(world: dict[str, Any]):
    db, data_dir = world["db"], world["data_dir"]
    templates = _templates()
    task = _task(db)
    evidence = _evidence(db, task, world["document_id"])
    first = generate(db, data_dir, REGISTRY_DIR, templates["approval-note"], GOOD_NOTE, task=task, author=ENGINEER, evidence=evidence)
    outcome = decide(db, data_dir, REGISTRY_DIR, list(templates.values()), artifact_id=first.artifact_id,
                     approver=APPROVER, approve=False, comment="State the re-inspection method.")
    assert outcome["decision"] == "rejected" and outcome["status"] == "VERIFIED"
    second = generate(
        db, data_dir, REGISTRY_DIR, templates["approval-note"],
        {**GOOD_NOTE, "recommendation": "Continue service; re-inspect CML-3 by UT within 12 months."},
        task=task, author=ENGINEER, evidence=evidence, revision_note="Revised: re-inspection method stated",
    )
    assert second.version == 2 and second.status == "VERIFIED"
    _, data = read_bytes(db, data_dir, second.artifact_id)
    text = rendered_text("docx", data)
    assert "rejected at review: State the re-inspection method." in text
    assert "Revised: re-inspection method stated" in text


@requires_pgvector
@pytest.mark.integration
def test_the_calculation_sheet_fills_by_column_heading_and_carries_its_markings(world: dict[str, Any]):
    db, data_dir = world["db"], world["data_dir"]
    template = _templates()["calculation-sheet"]
    task = _task(db)
    evidence = _evidence(db, task, world["document_id"])
    db.execute(
        "INSERT INTO task_evidence (task_id, evidence_id, kind, text, classification, detail) VALUES "
        "(%(t)s::uuid, 'C1', 'computation', '(9.2 - 8.4) / 0.25 = 3.2', 'internal', %(d)s)",
        {"t": task.task_id, "d": Json({"expression": "(9.2 - 8.4) / 0.25", "result": 3.2})},
    )
    evidence["C1"] = {"evidence_id": "C1", "kind": "computation", "text": "(9.2 - 8.4) / 0.25 = 3.2",
                      "classification": "internal", "detail": {"expression": "(9.2 - 8.4) / 0.25", "result": 3.2}}
    content = {
        "inputs": [
            {"quantity": "Measured minimum thickness", "value": "9.2", "unit": "mm", "citations": ["E1"]},
            {"quantity": "Minimum required thickness", "value": "8.4", "unit": "mm", "citations": ["E1"]},
        ],
        "working": [{"step": "Remaining life", "expression": "(9.2 - 8.4) / 0.25", "result": "3.2", "citations": ["C1"]}],
        "result": "3.2 years [C1]",
        "references": ["Inspection report IR-2026-0147 [E1]"],
    }
    generated = generate(db, data_dir, REGISTRY_DIR, template, content, task=task, author=ENGINEER, evidence=evidence)
    assert generated.verification.passed, generated.verification.to_dict()
    _, data = read_bytes(db, data_dir, generated.artifact_id)
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(data)).active
    assert sheet is not None
    assert [sheet.cell(row=13, column=c).value for c in range(1, 5)] == ["Measured minimum thickness", 9.2, "mm", "[1]"]
    assert [sheet.cell(row=22, column=c).value for c in range(1, 5)] == ["Remaining life", "(9.2 - 8.4) / 0.25", 3.2, "[2]"]
    assert sheet["E1"].value == "INTERNAL" and "INTERNAL" in (sheet.oddFooter.center.text or "")
    assert sheet.cell(row=48, column=2).value.startswith("Computation C1")


# -- a report shaped by its request ---------------------------------------------------------

GOOD_REPORT: dict[str, Any] = {
    "title": "Heat exchanger E-101 - shell condition",
    "summary": "E-101's thinnest reading leaves 3.2 years of remaining life [E1].",
    "body": [
        {"heading": "Shell condition",
         "text": "CML-3 measured 9.2 mm against a t-min of 8.4 mm [E1].\nThe shell is carbon steel, nominal 12.0 mm [E2]."},
        {"heading": "Corrosion and remaining life",
         "text": "- The corrosion rate is 0.25 mm/year [E1]\n- Remaining life is 3.2 years [E1]"},
    ],
    "recommendations": ["Re-inspect CML-3 within 12 months."],
    "open_questions": ["Is a shell repair planned before 2028?"],
}


def test_a_report_body_is_normalised_into_headed_sections_of_cited_paragraphs():
    template = _templates()["report"]
    content = normalise(template, GOOD_REPORT)
    assert not content.issues
    body = content.sections["body"]
    assert [i.heading for i in body.items] == ["Shell condition", "Corrosion and remaining life"]
    assert [p.citations for p in body.items[0].parts] == [["E1"], ["E2"]]
    assert body.items[1].parts[0].text == "- The corrosion rate is 0.25 mm/year"
    loose = normalise(template, {**GOOD_REPORT, "body": [{"text": "no heading [E1]"}, {"heading": "Empty"}, "just text"]})
    assert any("needs a heading" in i for i in loose.issues)
    assert any("'Empty' has no text" in i for i in loose.issues)
    assert any("expected an object with 'heading' and 'text'" in i for i in loose.issues)


@requires_pgvector
@pytest.mark.integration
def test_a_report_is_rendered_with_its_own_headings_and_references_where_they_belong(world: dict[str, Any]):
    db, data_dir = world["db"], world["data_dir"]
    template = _templates()["report"]
    task = _task(db)
    evidence = _evidence(db, task, world["document_id"])
    generated = generate(db, data_dir, REGISTRY_DIR, template, GOOD_REPORT, task=task, author=ENGINEER, evidence=evidence)
    tiers = {t.name: t.status for t in generated.verification.tiers}
    assert tiers == {"structural": "pass", "schema": "pass", "citation": "pass", "grounding": "pass"}, generated.verification.to_dict()
    artifact, data = read_bytes(db, data_dir, generated.artifact_id)
    assert generated.status == "VERIFIED" and not artifact["requires_approval"]  # the author owns a report
    text = rendered_text("docx", data)
    assert "Shell condition" in text and "Corrosion and remaining life" in text
    assert "CML-3 measured 9.2 mm against a t-min of 8.4 mm [1]." in text
    assert "The shell is carbon steel, nominal 12.0 mm [2]." in text
    assert "•  The corrosion rate is 0.25 mm/year [1]" in text
    assert "Re-inspect CML-3 within 12 months." in text and "{{" not in text
    assert "Recommendations" in text and "Open questions" in text
    assert "<h3>Shell condition</h3>" in to_html("docx", data)
    assert artifact["title"].startswith("Report: Heat exchanger E-101")
    # Asked for two things, the report is those two things: optional sections left empty
    # leave the document with their headings, rather than reading "Not applicable."
    only_body = {k: v for k, v in GOOD_REPORT.items() if k not in ("recommendations", "open_questions")}
    lean = generate(db, data_dir, REGISTRY_DIR, template, {**only_body, "title": "Report on E-101"}, task=task,
                    author=ENGINEER, evidence=evidence)
    assert lean.status == "VERIFIED", lean.verification.to_dict()
    lean_artifact, lean_data = read_bytes(db, data_dir, lean.artifact_id)
    lean_text = rendered_text("docx", lean_data)
    assert "Recommendations" not in lean_text and "Open questions" not in lean_text and "Not applicable" not in lean_text
    assert "Shell condition" in lean_text and "Sources" in lean_text and "Revision history" in lean_text
    assert lean_artifact["title"] == "Report on E-101"  # a subject that names its kind is not prefixed again


@requires_pgvector
@pytest.mark.integration
def test_a_report_section_must_cite_and_each_paragraph_is_grounded_on_its_own(world: dict[str, Any]):
    db, data_dir = world["db"], world["data_dir"]
    template = _templates()["report"]
    task = _task(db)
    evidence = _evidence(db, task, world["document_id"])
    uncited = {**GOOD_REPORT, "body": [*GOOD_REPORT["body"], {"heading": "Opinion", "text": "It looks fine."}]}
    failed = generate(db, data_dir, REGISTRY_DIR, template, uncited, task=task, author=ENGINEER, evidence=evidence)
    assert failed.status == "TEMP"
    citation = next(t for t in failed.verification.tiers if t.name == "citation")
    assert citation.status == "fail" and any("('Opinion') carries no citation" in i for i in citation.issues)
    # 14 mm is cited to E2, which says 12.0 mm: flagged -- even though E1 elsewhere in the
    # same section would not have saved it, each paragraph answers for its own numbers.
    wrong = {**GOOD_REPORT, "body": [{"heading": "Shell", "text": "CML-3 is 9.2 mm [E1].\nThe shell is 14 mm thick [E2]."}]}
    flagged = generate(db, data_dir, REGISTRY_DIR, template, wrong, task=task, author=ENGINEER, evidence=evidence)
    assert flagged.status == "VERIFIED"
    assert [c["number"] for c in flagged.verification.flagged_claims] == ["14"]
