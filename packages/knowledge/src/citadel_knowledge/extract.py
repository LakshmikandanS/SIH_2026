"""Getting text and page images out of an uploaded file, on the CPU.

`detect -> extract | OCR` from packages/knowledge/AGENTS.md. A born-digital PDF keeps
its text layer and word positions (pdfplumber); a page whose text layer is empty or
negligible is a scan and goes to Tesseract (`ocr.py`); an image is a one-page scan.
Every page is also rendered to an image, because citations are highlighted on the
page image and the vision model reads images, not text.
"""

from __future__ import annotations

import io
import threading
from pathlib import Path
from typing import Optional

from PIL import Image

from citadel_knowledge.normalise import Block, PageContent, clamp_bbox

#: A page with fewer extractable characters than this is treated as a scan.
_SCANNED_BELOW_CHARS = 40
RENDER_DPI = 170

PDF_TYPES = ("application/pdf",)
IMAGE_TYPES = ("image/png", "image/jpeg", "image/tiff", "image/bmp", "image/webp")
TEXT_TYPES = ("text/plain", "text/markdown", "text/csv")
DOCX_TYPES = ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",)


class UnsupportedDocument(ValueError):
    pass


def detect_kind(filename: str, head: bytes) -> tuple[str, str]:
    """(kind, mime_type) from the file's own magic bytes first, its name second."""
    name = filename.lower()
    if head.startswith(b"%PDF"):
        return "pdf", "application/pdf"
    if head.startswith(b"\x89PNG"):
        return "image", "image/png"
    if head[:3] == b"\xff\xd8\xff":
        return "image", "image/jpeg"
    if head[:4] in (b"II*\x00", b"MM\x00*"):
        return "image", "image/tiff"
    if head.startswith(b"PK") and name.endswith(".docx"):
        return "docx", DOCX_TYPES[0]
    if name.endswith((".txt", ".md", ".csv")):
        try:
            head.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise UnsupportedDocument(f"{filename}: not UTF-8 text") from exc
        return "text", "text/plain"
    raise UnsupportedDocument(f"{filename}: unsupported file type (PDF, PNG/JPEG/TIFF, DOCX or UTF-8 text)")


def _group_words(words: list[dict[str, float | str]], width: float, height: float) -> list[Block]:
    """Words -> lines -> blocks, by position: a new line when the baseline moves, a
    new block when the vertical gap is larger than a normal line spacing."""
    ordered = sorted(words, key=lambda w: (round(float(w["top"]) / 3), float(w["x0"])))
    lines: list[list[dict[str, float | str]]] = []
    for word in ordered:
        if lines and abs(float(lines[-1][0]["top"]) - float(word["top"])) <= 3.0:
            lines[-1].append(word)
        else:
            lines.append([word])
    blocks: list[Block] = []
    current: list[list[dict[str, float | str]]] = []

    def flush() -> None:
        if not current:
            return
        text = "\n".join(" ".join(str(w["text"]) for w in sorted(line, key=lambda w: float(w["x0"]))) for line in current)
        x0 = min(float(w["x0"]) for line in current for w in line)
        x1 = max(float(w["x1"]) for line in current for w in line)
        top = min(float(w["top"]) for line in current for w in line)
        bottom = max(float(w["bottom"]) for line in current for w in line)
        blocks.append(Block("text", "text_layer", text, clamp_bbox(x0 / width, top / height, x1 / width, bottom / height)))
        current.clear()

    previous_bottom: Optional[float] = None
    for line in lines:
        top = min(float(w["top"]) for w in line)
        bottom = max(float(w["bottom"]) for w in line)
        line_height = max(bottom - top, 1.0)
        if previous_bottom is not None and top - previous_bottom > 0.9 * line_height:
            flush()
        current.append(line)
        previous_bottom = bottom
    flush()
    return blocks


def pdf_pages(path: Path) -> list[PageContent]:
    """Text-layer blocks per page; pages that come back empty are marked for OCR."""
    import pdfplumber

    pages: list[PageContent] = []
    with pdfplumber.open(str(path)) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(keep_blank_chars=False, use_text_flow=False) or []
            blocks = _group_words(words, float(page.width), float(page.height)) if words else []
            chars = sum(len(b.text) for b in blocks)
            if chars < _SCANNED_BELOW_CHARS:
                pages.append(PageContent(page=index, text_source="none"))
            else:
                pages.append(PageContent(page=index, text_source="text_layer", blocks=blocks))
    return pages


#: PDFium is not thread-safe -- two threads rendering at once crash the process (seen
#: as a segfault in a worker running several ingestion threads). One lock per process
#: serialises the rendering only; OCR, vision and embedding still run concurrently.
_PDFIUM = threading.Lock()


def render_pdf(path: Path, dpi: int = RENDER_DPI) -> list[Image.Image]:
    import pypdfium2 as pdfium

    with _PDFIUM:
        document = pdfium.PdfDocument(str(path))
        try:
            return [document[i].render(scale=dpi / 72).to_pil().convert("RGB") for i in range(len(document))]
        finally:
            document.close()


def load_image(path: Path) -> Image.Image:
    with Image.open(path) as image:
        image.load()
        return image.convert("RGB")


def text_pages(path: Path) -> list[PageContent]:
    text = path.read_text(encoding="utf-8")
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    return [PageContent(page=1, text_source="text_layer", blocks=[Block("text", "text_layer", p) for p in paragraphs])]


def docx_pages(path: Path) -> list[PageContent]:
    from docx import Document

    document = Document(str(path))
    blocks = [Block("text", "text_layer", p.text.strip()) for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            blocks.append(Block("table", "text_layer", "\n".join(rows)))
    return [PageContent(page=1, text_source="text_layer", blocks=blocks)]


def to_jpeg(image: Image.Image, *, max_side: Optional[int] = None, quality: int = 82) -> bytes:
    if max_side and max(image.size) > max_side:
        image = image.copy()
        image.thumbnail((max_side, max_side))
    buffer = io.BytesIO()
    image.save(buffer, "JPEG", quality=quality, optimize=True)
    return buffer.getvalue()


__all__ = [
    "UnsupportedDocument",
    "detect_kind",
    "pdf_pages",
    "render_pdf",
    "load_image",
    "text_pages",
    "docx_pages",
    "to_jpeg",
    "RENDER_DPI",
]
