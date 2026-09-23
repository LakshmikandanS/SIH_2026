"""sheet.read and sheet.write: spreadsheets in the task workspace (.xlsx or .csv).

`cells` for sheet.write is a mapping of A1-style references to values
(`{"A1": "Tag", "B1": "E-101", "B2": "=B1"}`); a `rows` key holding a list of rows
writes a block starting at A1. Formulas are stored as written -- Excel computes them on
open -- and nothing here evaluates a formula on the model's behalf (that is
calc.evaluate's job, with its working kept).
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any

from citadel_contracts.domain import Resource

from citadel_tools.context import Invocation, ResourceNotFound, ToolContext, ToolFailure, ToolOutput
from citadel_tools.fs import normalised, workspace_path
from citadel_tools.plugins import ToolPlugin, task_resource

MAX_ROWS = 60
MAX_COLS = 20
MAX_WRITES = 2000
_REF = re.compile(r"^([A-Za-z]{1,3})([1-9][0-9]{0,5})$")


def _resource(ctx: ToolContext, args: dict[str, Any]) -> Resource:
    return task_resource(ctx, "workspace", normalised(ctx, str(args["path"])))


def _column_index(letters: str) -> int:
    index = 0
    for char in letters.upper():
        index = index * 26 + (ord(char) - 64)
    return index


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value if isinstance(value, (int, float, str)) else str(value)


def _read_csv(path: Path) -> list[list[Any]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return [row for _, row in zip(range(MAX_ROWS), csv.reader(handle))]


def _read(ctx: ToolContext, args: dict[str, Any], invocation: Invocation) -> ToolOutput:
    path = workspace_path(ctx, str(args["path"]))
    if not path.is_file():
        raise ResourceNotFound(f"{args['path']} does not exist in the task workspace")
    relative = path.relative_to(ctx.workspace().resolve()).as_posix()
    if path.suffix.lower() == ".csv":
        rows = [[_cell(v) for v in row[:MAX_COLS]] for row in _read_csv(path)]
        return ToolOutput({"path": relative, "sheet": "csv", "rows": rows}, f"read {relative} ({len(rows)} rows)")
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ToolFailure("sheet.read reads .xlsx or .csv files")
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=False, read_only=True)
    name = str(args.get("sheet") or workbook.sheetnames[0])
    if name not in workbook.sheetnames:
        raise ResourceNotFound(f"no sheet {name!r}; sheets are {workbook.sheetnames}")
    sheet = workbook[name]
    rows = [
        [_cell(v) for v in row[:MAX_COLS]]
        for row in sheet.iter_rows(max_row=MAX_ROWS, values_only=True)
    ]
    while rows and not any(v != "" for v in rows[-1]):
        rows.pop()
    return ToolOutput(
        {"path": relative, "sheet": name, "sheets": workbook.sheetnames, "rows": rows,
         "truncated": (sheet.max_row or 0) > MAX_ROWS},
        f"read {relative} [{name}] ({len(rows)} rows)",
    )


def _writes(cells: dict[str, Any]) -> list[tuple[int, int, Any]]:
    writes: list[tuple[int, int, Any]] = []
    for key, value in cells.items():
        if key == "rows":
            if not isinstance(value, list):
                raise ToolFailure("'rows' must be a list of rows")
            for r, row in enumerate(value, start=1):
                for c, item in enumerate(row if isinstance(row, list) else [row], start=1):
                    writes.append((r, c, item))
            continue
        match = _REF.match(key.strip())
        if not match:
            raise ToolFailure(f"{key!r} is not a cell reference like A1")
        writes.append((int(match.group(2)), _column_index(match.group(1)), value))
    if len(writes) > MAX_WRITES:
        raise ToolFailure(f"at most {MAX_WRITES} cells per call")
    return writes


def _write(ctx: ToolContext, args: dict[str, Any], invocation: Invocation) -> ToolOutput:
    path = workspace_path(ctx, str(args["path"]))
    writes = _writes(dict(args["cells"]))
    relative = normalised(ctx, str(args["path"]))
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".csv":
        grid = _read_csv(path) if path.exists() else []
        for r, c, value in writes:
            while len(grid) < r:
                grid.append([])
            row = grid[r - 1]
            while len(row) < c:
                row.append("")
            row[c - 1] = "" if value is None else str(value)
        buffer = io.StringIO()
        csv.writer(buffer).writerows(grid)
        path.write_text(buffer.getvalue(), encoding="utf-8")
        return ToolOutput({"path": relative, "cells_written": len(writes)}, f"wrote {len(writes)} cell(s) to {relative}")
    if path.suffix.lower() != ".xlsx":
        raise ToolFailure("sheet.write writes .xlsx or .csv files")
    from openpyxl import Workbook, load_workbook

    workbook = load_workbook(str(path)) if path.exists() else Workbook()
    name = str(args.get("sheet") or (workbook.sheetnames[0] if path.exists() else "Sheet1"))
    if name in workbook.sheetnames:
        sheet = workbook[name]
    elif not path.exists():
        sheet = workbook.active
        sheet.title = name
    else:
        sheet = workbook.create_sheet(name)
    for r, c, value in writes:
        sheet.cell(row=r, column=c, value=value if isinstance(value, (int, float, str)) or value is None else str(value))
    workbook.save(str(path))
    return ToolOutput({"path": relative, "sheet": name, "cells_written": len(writes)},
                      f"wrote {len(writes)} cell(s) to {relative} [{name}]")


PLUGINS = {
    "sheet.read": ToolPlugin("sheet.read", _resource, _read),
    "sheet.write": ToolPlugin("sheet.write", _resource, _write),
}

__all__ = ["PLUGINS"]
